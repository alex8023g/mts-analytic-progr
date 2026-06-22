import io
from collections.abc import Callable, Iterable
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import Division, Employee

COLUMNS: list[tuple[str, int]] = [
    ("ФИО", 32),
    ("Должность", 28),
    ("Департамент", 28),
    ("Отдел", 24),
    ("Руководитель", 28),
    ("Дата приема", 14),
    ("Дата увольнения", 16),
    ("Статус", 12),
    ("Штат", 16),
    ("Зарплата", 14),
]
DIVISION_COLUMNS: list[tuple[str, int]] = [
    ("Отдел", 40),
]
DATE_FORMAT = "DD.MM.YYYY"
MONEY_FORMAT = "#,##0 ₽"


def _status(employee: Employee, ref_date: date) -> str:
    fired_at = employee.fired_at
    if fired_at is not None and fired_at <= ref_date:
        return "Уволен"
    return "Работает"


def _write_header(ws: Worksheet, columns: list[tuple[str, int]]) -> None:
    header_font = Font(bold=True)
    for idx, (title, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.font = header_font
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"


def _write_employees_sheet(
    ws: Worksheet,
    employees: Iterable[Employee],
    ref_date: date,
    progress: Callable[[int, int], None] | None = None,
) -> None:
    ws.title = "Сотрудники"
    _write_header(ws, COLUMNS)

    employees = list(employees)
    total = len(employees)
    for processed, employee in enumerate(employees, start=1):
        row = [
            employee.full_name,
            employee.position,
            employee.department,
            employee.division.name if employee.division else None,
            employee.manager,
            employee.hired_at,
            employee.fired_at,
            _status(employee, ref_date),
            employee.staff_type,
            employee.salary,
        ]
        ws.append(row)
        excel_row = ws.max_row
        ws.cell(row=excel_row, column=6).number_format = DATE_FORMAT
        ws.cell(row=excel_row, column=7).number_format = DATE_FORMAT
        ws.cell(row=excel_row, column=10).number_format = MONEY_FORMAT
        if progress is not None:
            progress(processed, total)


def _write_divisions_sheet(
    ws: Worksheet,
    divisions: Iterable[Division],
    progress: Callable[[int, int], None] | None = None,
) -> None:
    ws.title = "Отделы"
    _write_header(ws, DIVISION_COLUMNS)

    divisions = list(divisions)
    total = len(divisions)
    for processed, division in enumerate(divisions, start=1):
        ws.append([division.name])
        if progress is not None:
            progress(processed, total)


def _to_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_employees_xlsx(
    employees: Iterable[Employee],
    relevance_date: date | None = None,
    progress: Callable[[int, int], None] | None = None,
) -> bytes:
    ref_date = relevance_date or date.today()

    wb = Workbook()
    _write_employees_sheet(wb.active, employees, ref_date, progress)
    return _to_bytes(wb)


def build_divisions_xlsx(
    divisions: Iterable[Division],
    progress: Callable[[int, int], None] | None = None,
) -> bytes:
    wb = Workbook()
    _write_divisions_sheet(wb.active, divisions, progress)
    return _to_bytes(wb)


def build_employees_and_divisions_xlsx(
    employees: Iterable[Employee],
    divisions: Iterable[Division],
    relevance_date: date | None = None,
    progress: Callable[[int, int], None] | None = None,
) -> bytes:
    ref_date = relevance_date or date.today()

    wb = Workbook()
    # Progress tracks the (dominant) employees sheet; divisions are appended after.
    _write_employees_sheet(wb.active, employees, ref_date, progress)
    _write_divisions_sheet(wb.create_sheet(), divisions)
    return _to_bytes(wb)
